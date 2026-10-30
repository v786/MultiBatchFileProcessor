import zipfile
from typing import IO

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from pipeline.errors.model import CorruptFile
from pipeline.extract import MAX_ROWS, Extracted


def extract(stream: IO[bytes], out: Extracted) -> None:
    try:
        wb = load_workbook(stream, read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise CorruptFile("not a valid XLSX workbook") from exc
    sheets = []
    try:
        for ws in wb.worksheets:
            rows_iter = ws.iter_rows(values_only=True)
            header = next(rows_iter, None)
            if header is None:
                sheets.append({"name": ws.title, "columns": [], "rowCount": 0, "rows": []})
                continue
            columns = [str(c) if c is not None else f"col{i + 1}" for i, c in enumerate(header)]
            rows, count = [], 0
            for row in rows_iter:
                if all(v is None for v in row):
                    continue
                count += 1
                if len(rows) < MAX_ROWS:
                    rows.append({columns[i]: v for i, v in enumerate(row) if i < len(columns)})
            if count > MAX_ROWS:
                out.warn(f"sheet {ws.title!r}: only the first {MAX_ROWS} of {count} rows are included")
            sheets.append({"name": ws.title, "columns": columns, "rowCount": count, "rows": rows})
    finally:
        wb.close()
    out.data = {"sheets": sheets}
